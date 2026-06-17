import os
import re
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, color_hex):
    """Set the background color of a table cell."""
    shading_elm = OxmlElement('w:shd')
    shading_elm.set(qn('w:val'), 'clear')
    shading_elm.set(qn('w:color'), 'auto')
    shading_elm.set(qn('w:fill'), color_hex)
    cell._tc.get_or_add_tcPr().append(shading_elm)

def add_formatted_text(p, text):
    """Add text to a paragraph, parsing **bold** style."""
    parts = text.split('**')
    is_bold = False
    for part in parts:
        if not part:
            is_bold = not is_bold
            continue
        run = p.add_run(part)
        run.bold = is_bold
        is_bold = not is_bold

def parse_markdown_to_docx(md_path, docx_path):
    doc = Document()
    paper_dir = os.path.dirname(os.path.abspath(md_path))
    
    # Page Setup: Standard 1-inch margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        
    # Styles Setup
    style_normal = doc.styles['Normal']
    font = style_normal.font
    font.name = 'Times New Roman'
    font.size = Pt(12)
    style_normal.paragraph_format.line_spacing = 1.15
    style_normal.paragraph_format.space_after = Pt(6)
    
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()
        
    in_table = False
    table_rows = []
    
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        
        # Table parsing
        if line.strip().startswith('|'):
            in_table = True
            table_rows.append(line)
            i += 1
            continue
        elif in_table:
            # End of table, process it
            in_table = False
            if table_rows:
                # Remove separator row like |---|---|
                cleaned_rows = []
                for r in table_rows:
                    if re.match(r'^\s*\|\s*[-:]+\s*\|', r) or all(c in '|-: \t' for c in r.strip()):
                        continue
                    cleaned_rows.append(r)
                
                if cleaned_rows:
                    parsed_table = []
                    for r in cleaned_rows:
                        cells = [c.strip() for c in r.split('|')[1:-1]]
                        parsed_table.append(cells)
                    
                    if parsed_table:
                        num_rows = len(parsed_table)
                        num_cols = len(parsed_table[0])
                        table = doc.add_table(rows=num_rows, cols=num_cols)
                        table.style = 'Table Grid'
                        
                        for row_idx, row_data in enumerate(parsed_table):
                            row_cells = table.rows[row_idx].cells
                            for col_idx in range(min(num_cols, len(row_data))):
                                cell = row_cells[col_idx]
                                cell.text = ""
                                p = cell.paragraphs[0]
                                p.paragraph_format.space_after = Pt(2)
                                p.paragraph_format.line_spacing = 1.0
                                run = p.add_run(row_data[col_idx])
                                if row_idx == 0:
                                    run.bold = True
                                    set_cell_background(cell, 'F2F2F2')
                        
                        doc.add_paragraph() # space after table
            table_rows = []
            
        # Image parsing
        match_img = re.match(r'^!\[(.*?)\]\((.*?)\)', line.strip())
        if match_img:
            caption = match_img.group(1)
            img_name = match_img.group(2)
            img_path = os.path.join(paper_dir, img_name)
            if os.path.exists(img_path):
                p_img = doc.add_paragraph()
                p_img.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_img.paragraph_format.space_before = Pt(12)
                p_img.paragraph_format.space_after = Pt(6)
                r_img = p_img.add_run()
                r_img.add_picture(img_path, width=Inches(5.8))
                
                p_cap = doc.add_paragraph()
                p_cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_cap.paragraph_format.space_after = Pt(12)
                r_cap = p_cap.add_run(caption)
                r_cap.font.name = 'Times New Roman'
                r_cap.font.size = Pt(10)
                r_cap.font.italic = True
            else:
                print(f"Warning: Image file not found: {img_path}")
            i += 1
            continue
            
        # Headers
        if line.startswith('# '):
            title_text = line[2:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(12)
            run = p.add_run(title_text)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(18)
            run.bold = True
            
            # Author placeholder block
            p_auth = doc.add_paragraph()
            p_auth.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_auth.paragraph_format.space_after = Pt(24)
            run_auth = p_auth.add_run("Anonymous Authors for Double-Blind Peer Review")
            run_auth.font.italic = True
            run_auth.font.size = Pt(11)
            
        elif line.startswith('## '):
            heading_text = line[3:].strip()
            p = doc.add_paragraph()
            run = p.add_run(heading_text)
            run.font.name = 'Times New Roman'
            run.bold = True
            if heading_text.lower() == 'abstract':
                p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after = Pt(12)
                run.font.size = Pt(14)
            else:
                p.paragraph_format.space_before = Pt(12)
                p.paragraph_format.space_after = Pt(6)
                run.font.size = Pt(14)
                
        elif line.startswith('### '):
            heading_text = line[4:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(heading_text)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(12)
            run.bold = True
            
        # Lists
        elif line.strip().startswith('* ') or line.strip().startswith('- '):
            content = line.strip()[2:]
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_after = Pt(4)
            add_formatted_text(p, content)
            
        elif re.match(r'^\d+\.\s', line.strip()):
            content = re.sub(r'^\d+\.\s', '', line.strip())
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.space_after = Pt(4)
            num_prefix = re.match(r'^(\d+\.\s)', line.strip()).group(1)
            p.add_run(num_prefix).bold = True
            add_formatted_text(p, content)
            
        # Horizontal rule
        elif line.strip() == '---':
            pass
            
        # Empty lines
        elif not line.strip():
            pass
            
        # Paragraph text
        else:
            p = doc.add_paragraph()
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.space_after = Pt(8)
            add_formatted_text(p, line)
            
        i += 1
        
    doc.save(docx_path)
    print(f"Compiled paper to {docx_path}")

def run():
    paper_dir = os.path.dirname(os.path.abspath(__file__))
    md_file = os.path.join(paper_dir, 'driftguard_paper.md')
    
    file1 = os.path.join(paper_dir, 'DriftGuard_Research_Paper.docx')
    file2 = os.path.join(paper_dir, 'DriftGuard_Research_Paper_Extended.docx')
    file3 = os.path.join(paper_dir, 'DriftGuard_Research_Paper_v3.docx')
    
    parse_markdown_to_docx(md_file, file1)
    parse_markdown_to_docx(md_file, file2)
    parse_markdown_to_docx(md_file, file3)
    print("All paper Word documents compiled successfully.")

if __name__ == "__main__":
    run()
